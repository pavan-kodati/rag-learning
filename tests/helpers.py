"""Builders for test documents."""

from __future__ import annotations

import io

from reportlab.lib.pagesizes import A4
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.pdfgen import canvas

from legal_workspace.ingestion.pdfio import parse_pdf
from legal_workspace.ingestion.render import text_to_pdf


def pdf(text: str) -> bytes:
    return text_to_pdf(text)


def pages_of(text: str):  # noqa: ANN201
    return parse_pdf(text_to_pdf(text))[0]


def encrypted_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, encrypt=StandardEncryption("secret", canPrint=0))
    c.drawString(72, 700, "secret text that nobody can read")
    c.save()
    return buf.getvalue()


def scanned_pdf(text: str) -> bytes:
    """An image-only PDF (no text layer) made by rasterising a rendered page."""
    import pypdfium2 as pdfium

    page = pdfium.PdfDocument(text_to_pdf(text, footer=False))[0]
    img = page.render(scale=2.5).to_pil().convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="PDF", resolution=180.0)
    return buf.getvalue()


def docx_bytes(paragraphs: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def png_bytes(text: str = "Hello") -> bytes:
    import pypdfium2 as pdfium

    page = pdfium.PdfDocument(text_to_pdf(text, footer=False))[0]
    buf = io.BytesIO()
    page.render(scale=2).to_pil().convert("RGB").save(buf, format="PNG")
    return buf.getvalue()


class DocScript:
    """A scripted 'model' for integration tests: answers classification/extraction by a marker found in the prompt."""

    def __init__(self) -> None:
        self.classify: dict[str, dict] = {}
        self.extract: dict[str, list[dict]] = {}
        self.fail_for: set[str] = set()
        self.calls = 0

    def add(self, marker: str, classification: dict, items: list[dict] | None = None) -> None:
        self.classify[marker] = classification
        self.extract[marker] = items or []

    def __call__(self, task, system, user, schema, meta):  # noqa: ANN001, ANN204
        from legal_workspace.errors import ModelUnavailable

        self.calls += 1
        for marker in self.fail_for:
            if marker in user:
                raise ModelUnavailable("scripted outage")
        for marker, payload in (self.classify if task == "classify" else self.extract).items():
            if marker in user:
                return payload if task == "classify" else {"items": payload}
        return (
            {"document_type": "LETTER", "title": "unknown", "document_date": "", "confidence": 0.2}
            if task == "classify"
            else {"items": []}
        )


def cls(
    doc_type: str,
    title: str,
    date: str,
    *,
    confidence: float = 0.95,
    author: str = "",
    parties=None,
    signed=False,
    amends="",
    terms=None,
) -> dict:
    return {
        "document_type": doc_type,
        "title": title,
        "document_date": date,
        "author": author,
        "parties": parties or [],
        "defined_terms": terms or [],
        "signed_by_all_parties": signed,
        "amends": amends,
        "confidence": confidence,
    }
