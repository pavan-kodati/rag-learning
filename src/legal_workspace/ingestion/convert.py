"""Anything -> PDF, so every document has pages to cite (blueprint §7.2)."""

from __future__ import annotations

import email
import html
import io
import logging
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from email import policy
from pathlib import Path

from ..errors import ParseError
from .pdfio import open_pdf
from .render import text_to_pdf
from .validate import UploadInfo, sanitize_filename

log = logging.getLogger(__name__)
MAX_ATTACHMENTS = 10
_QUOTE_INTRO = re.compile(r"(?im)^\s*(on .{5,120}wrote:|-{2,}\s*original message\s*-{2,}|from:.*sent:.*)\s*$")


@dataclass
class Converted:
    pdf: bytes
    attachments: list[tuple[str, bytes]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def strip_quoted_replies(body: str) -> str:
    """Drop '> quoted' lines and everything after an 'On ... wrote:' / 'Original Message' marker."""
    out: list[str] = []
    for line in body.splitlines():
        if _QUOTE_INTRO.match(line):
            break
        if line.lstrip().startswith(">"):
            continue
        out.append(line)
    return "\n".join(out).strip()


def _html_to_text(markup: str) -> str:
    markup = re.sub(r"(?is)<(script|style).*?</\1>", " ", markup)
    markup = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", markup)
    return html.unescape(re.sub(r"<[^>]+>", " ", markup))


def _eml_to_text(data: bytes) -> tuple[str, list[tuple[str, bytes]]]:
    try:
        msg = email.message_from_bytes(data, policy=policy.default)
    except Exception as e:  # noqa: BLE001
        raise ParseError("the e-mail could not be parsed") from e
    headers = [f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h]]
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if body_part is not None:
        content = body_part.get_content()
        body = _html_to_text(content) if body_part.get_content_type() == "text/html" else content
    attachments: list[tuple[str, bytes]] = []
    for part in msg.iter_attachments():
        payload = part.get_payload(decode=True)
        fname = part.get_filename()
        if payload and fname and len(attachments) < MAX_ATTACHMENTS:
            attachments.append((sanitize_filename(fname), bytes(payload)))
    return "\n".join(headers) + "\n\n" + strip_quoted_replies(body), attachments


def _docx_via_libreoffice(data: bytes) -> bytes | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "in.docx"
        src.write_bytes(data)
        try:
            subprocess.run(  # noqa: S603
                [exe, "--headless", "--convert-to", "pdf", "--outdir", d, str(src)],
                check=True, capture_output=True, timeout=120, env={"HOME": d, "PATH": "/usr/bin:/bin:/usr/local/bin"},
            )  # fmt: skip
        except (subprocess.SubprocessError, OSError) as e:
            log.warning("LibreOffice conversion failed: %s", e)
            return None
        out = Path(d) / "in.pdf"
        return out.read_bytes() if out.exists() else None


def _docx_to_text(data: bytes) -> str:
    try:
        import docx

        doc = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ParseError("the Word document could not be read (corrupt or password-protected)") from e
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _image_to_pdf(data: bytes) -> bytes:
    from PIL import Image, ImageSequence

    try:
        img = Image.open(io.BytesIO(data))
        frames = [f.convert("RGB") for f in ImageSequence.Iterator(img)][:200]
    except Exception as e:  # noqa: BLE001
        raise ParseError("the image could not be read") from e
    buf = io.BytesIO()
    frames[0].save(buf, format="PDF", resolution=200.0, save_all=True, append_images=frames[1:])
    return buf.getvalue()


def to_pdf(info: UploadInfo, data: bytes) -> Converted:
    if info.kind == "pdf":
        open_pdf(data)  # validates: corrupt / encrypted / empty / too many pages
        return Converted(data)
    if info.kind == "docx":
        pdf = _docx_via_libreoffice(data)
        if pdf:
            open_pdf(pdf)
            return Converted(pdf)
        text = _docx_to_text(data)
        if not text.strip():
            raise ParseError("the Word document contains no text")
        return Converted(
            text_to_pdf(text, title=info.clean_name),
            notes=["converted without LibreOffice (layout simplified)"],
        )
    if info.kind == "eml":
        text, attachments = _eml_to_text(data)
        if not text.strip():
            raise ParseError("the e-mail is empty")
        return Converted(text_to_pdf(text, title=info.clean_name), attachments)
    if info.kind == "image":
        return Converted(_image_to_pdf(data), notes=["scanned image: text comes from OCR"])
    text = _decode(data)
    if not text.strip():
        raise ParseError("the text file is empty")
    return Converted(text_to_pdf(text, title=info.clean_name))
